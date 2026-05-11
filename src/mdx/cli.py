from __future__ import annotations

import fnmatch
import shutil
import sys
from pathlib import Path

from docx import Document
from registers import CommandRegistry
from registers.cli import types as t

from . import __version__
from .config import load_config, write_default_config
from .converter import convert_markdown_to_docx, render_markdown_to_html
from .styles import STYLE_PRESETS
from .validation import ValidationReport, summarize_reports, validate_markdown_file

EXIT_OK = 0
EXIT_GENERAL = 1
EXIT_ARGS = 2
EXIT_VALIDATION = 3
EXIT_INPUT_NOT_FOUND = 4
EXIT_OUTPUT_PATH = 5
EXIT_CONFIG_OR_TEMPLATE = 6
EXIT_BACKEND = 7

registry = CommandRegistry()


def _error(message: str, *, code: int, path: Path | None = None, hint: str | None = None) -> None:
    print(f"Error: {message}", file=sys.stderr)
    if path is not None:
        print("\nPath:", file=sys.stderr)
        print(path, file=sys.stderr)
    if hint:
        print(f"\n{hint}", file=sys.stderr)
    raise SystemExit(code)


def _log(message: str, *, quiet: bool = False) -> None:
    if not quiet:
        print(message)


def _resolve_config(config: str, *, explicit: bool) -> dict:
    config_path = Path(config)
    try:
        return load_config(config_path, explicit=explicit)
    except FileNotFoundError:
        _error(
            "configuration file not found.",
            code=EXIT_CONFIG_OR_TEMPLATE,
            path=config_path,
            hint="Use --config with a valid path or run `mdx init`.",
        )
    except ValueError as exc:
        _error(str(exc), code=EXIT_CONFIG_OR_TEMPLATE, path=config_path)
    return {}


def _ensure_input_file(path: Path) -> None:
    if not path.exists() or not path.is_file():
        _error("input file not found.", code=EXIT_INPUT_NOT_FOUND, path=path)


def _ensure_input_dir(path: Path) -> None:
    if not path.exists() or not path.is_dir():
        _error("input directory not found.", code=EXIT_INPUT_NOT_FOUND, path=path)


def _ensure_write_target(path: Path, *, force: bool) -> None:
    if path.exists() and not force:
        _error(
            "output file already exists.",
            code=EXIT_OUTPUT_PATH,
            path=path,
            hint="Use --force to overwrite the file.",
        )
    path.parent.mkdir(parents=True, exist_ok=True)


def _resolve_style(style: str | None, config: dict) -> str:
    resolved = (style or config["style"]["preset"] or "professional").strip().lower()
    if resolved not in STYLE_PRESETS:
        _error(
            f"unsupported style '{resolved}'.",
            code=EXIT_ARGS,
            hint="Use one of: professional, technical, executive, academic, minimal.",
        )
    return resolved


def _resolve_template(template: str | None, config: dict) -> Path | None:
    raw = (template or config["style"].get("template") or "").strip()
    if not raw:
        return None
    candidate = Path(raw)
    if not candidate.exists():
        _error(
            "template file not found.",
            code=EXIT_CONFIG_OR_TEMPLATE,
            path=candidate,
            hint="Use --template with a valid .docx template path.",
        )
    return candidate


def _discover_markdown_files(input_dir: Path, include: str, exclude: str | None) -> list[Path]:
    files = [path for path in input_dir.rglob("*.md") if path.is_file()]
    selected: list[Path] = []
    for path in files:
        rel = path.relative_to(input_dir).as_posix()
        include_ok = fnmatch.fnmatch(rel, include) or fnmatch.fnmatch(path.name, include)
        exclude_hit = False
        if exclude:
            exclude_hit = fnmatch.fnmatch(rel, exclude) or fnmatch.fnmatch(path.name, exclude)
        if include_ok and not exclude_hit:
            selected.append(path)
    return sorted(selected)


def _print_report(report: ValidationReport, *, quiet: bool, verbose: bool) -> None:
    if quiet:
        return
    print(f"Validation report: {report.path}")
    if not report.issues:
        print("\nStatus: passed")
        return
    if verbose:
        print()
        for issue in report.issues:
            prefix = f"Line {issue.line}: " if issue.line else ""
            print(f"- {issue.severity.upper()}: {prefix}{issue.message}")
    else:
        warnings = report.warnings
        errors = report.errors
        print(f"\nWarnings: {warnings}")
        print(f"Errors:   {errors}")
    status = "passed" if report.passed else "errors found"
    print(f"\nStatus: {status}")


@registry.register(
    "init",
    description="Initialize an MDX project.",
    examples=["mdx init --with-template"],
    render=False,
)
@registry.argument("with_template", type=bool, default=False, help="Generate a default template.")
@registry.argument("force", type=bool, default=False, help="Overwrite existing files.")
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def init(with_template: bool = False, force: bool = False, quiet: bool = False) -> None:
    created: list[Path] = []
    config_path = Path("mdx.toml")
    template_dir = Path("templates")
    template_path = template_dir / "professional.docx"
    docs_dir = Path("docs")

    if not config_path.exists() or force:
        try:
            write_default_config(config_path, overwrite=force)
            created.append(config_path)
        except FileExistsError:
            pass

    template_dir.mkdir(parents=True, exist_ok=True)
    docs_dir.mkdir(parents=True, exist_ok=True)

    if with_template and (force or not template_path.exists()):
        try:
            doc = Document()
            doc.save(template_path)
            created.append(template_path)
        except Exception as exc:  # pragma: no cover - backend/runtime failure
            _error(f"failed to create template: {exc}", code=EXIT_BACKEND, path=template_path)

    created.append(docs_dir)

    if not quiet:
        print("MDX project initialized.\n")
        print("Created:")
        for path in created:
            print(f"- {path.as_posix()}")
        print("\nStatus: ready")
    return None


@registry.register(
    "convert",
    description="Convert one Markdown file into DOCX.",
    examples=["mdx convert docs/example.md --output docs/example.docx"],
    render=False,
)
@registry.argument("input_md", type=t.Path(exists=True, readable=True), help="Input Markdown file.")
@registry.argument("output", type=str, help="Output DOCX path.")
@registry.argument("template", type=str, default="", help="DOCX template path.")
@registry.argument("style", type=str, default="", help="Style preset.")
@registry.argument("config", type=str, default="mdx.toml", help="Config file path.")
@registry.argument("force", type=bool, default=False, help="Overwrite output if it exists.")
@registry.argument("strict", type=bool, default=False, help="Treat warnings as errors.")
@registry.argument("dry_run", type=bool, default=False, help="Show actions without writing.")
@registry.argument("verbose", type=bool, default=False, help="Enable detailed logs.")
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def convert(
    input_md: Path,
    output: str,
    template: str = "",
    style: str = "",
    config: str = "mdx.toml",
    force: bool = False,
    strict: bool = False,
    dry_run: bool = False,
    verbose: bool = False,
    quiet: bool = False,
) -> None:
    input_path = Path(input_md)
    output_path = Path(output)
    _ensure_input_file(input_path)
    _ensure_write_target(output_path, force=force)

    config_dict = _resolve_config(config, explicit=(config != "mdx.toml"))
    resolved_style = _resolve_style(style, config_dict)
    resolved_template = _resolve_template(template, config_dict)
    validation_cfg = config_dict.get("validation", {})

    report = validate_markdown_file(
        input_path,
        strict=False,
        check_pii=bool(validation_cfg.get("check_pii", False)),
        check_placeholders=bool(validation_cfg.get("check_placeholders", True)),
        check_tables=bool(validation_cfg.get("check_empty_tables", True)),
        check_headings=bool(validation_cfg.get("check_heading_order", True)),
        check_links=bool(validation_cfg.get("check_links", False)),
    )

    strict_mode = strict or bool(validation_cfg.get("strict", False))
    if report.errors or (strict_mode and report.warnings):
        _print_report(report, quiet=quiet, verbose=True)
        _error("validation failed.", code=EXIT_VALIDATION)

    if dry_run:
        _log("Dry run: conversion planned.", quiet=quiet)
        _log(f"  input:  {input_path.as_posix()}", quiet=quiet)
        _log(f"  output: {output_path.as_posix()}", quiet=quiet)
        _log(f"  style:  {resolved_style}", quiet=quiet)
        return None

    if verbose and not quiet:
        print(f"[mdx] reading input: {input_path.as_posix()}")
        print("[mdx] validating markdown")
        print(f"[mdx] applying style preset: {resolved_style}")
        if resolved_template:
            print(f"[mdx] using template: {resolved_template.as_posix()}")
        print(f"[mdx] writing output: {output_path.as_posix()}")

    try:
        convert_markdown_to_docx(
            input_path,
            output_path,
            style=resolved_style,
            template_path=resolved_template,
        )
    except Exception as exc:  # pragma: no cover - backend/runtime failure
        _error(f"conversion backend error: {exc}", code=EXIT_BACKEND)

    _log("Converting:", quiet=quiet)
    _log(f"  input:  {input_path.as_posix()}", quiet=quiet)
    _log(f"  output: {output_path.as_posix()}", quiet=quiet)
    _log(f"  style:  {resolved_style}", quiet=quiet)
    _log("\nValidation: passed", quiet=quiet)
    _log("DOCX written successfully.", quiet=quiet)
    return None


@registry.register(
    "batch",
    description="Convert a directory of Markdown files into DOCX.",
    examples=["mdx batch docs --output_dir docs --style professional --force"],
    render=False,
)
@registry.argument("input_dir", type=t.Path(exists=True), help="Directory containing Markdown files.")
@registry.argument("output_dir", type=str, help="Output directory.")
@registry.argument("include", type=str, default="*.md", help="Include glob pattern.")
@registry.argument("exclude", type=str, default="", help="Exclude glob pattern.")
@registry.argument("template", type=str, default="", help="DOCX template path.")
@registry.argument("style", type=str, default="", help="Style preset.")
@registry.argument("config", type=str, default="mdx.toml", help="Config file path.")
@registry.argument("force", type=bool, default=False, help="Overwrite output files.")
@registry.argument("strict", type=bool, default=False, help="Treat warnings as errors.")
@registry.argument("dry_run", type=bool, default=False, help="Show planned actions.")
@registry.argument("verbose", type=bool, default=False, help="Enable detailed logs.")
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def batch(
    input_dir: Path,
    output_dir: str,
    include: str = "*.md",
    exclude: str = "",
    template: str = "",
    style: str = "",
    config: str = "mdx.toml",
    force: bool = False,
    strict: bool = False,
    dry_run: bool = False,
    verbose: bool = False,
    quiet: bool = False,
) -> None:
    input_root = Path(input_dir)
    output_root = Path(output_dir)
    _ensure_input_dir(input_root)

    config_dict = _resolve_config(config, explicit=(config != "mdx.toml"))
    resolved_style = _resolve_style(style, config_dict)
    resolved_template = _resolve_template(template, config_dict)
    validation_cfg = config_dict.get("validation", {})
    strict_mode = strict or bool(validation_cfg.get("strict", False))

    files = _discover_markdown_files(input_root, include=include, exclude=exclude or None)
    if not files:
        _error("no markdown files matched include/exclude rules.", code=EXIT_INPUT_NOT_FOUND)

    converted = 0
    reports: list[ValidationReport] = []
    _log("Batch conversion started.\n", quiet=quiet)
    _log(f"Input directory:  {input_root.as_posix()}", quiet=quiet)
    _log(f"Output directory: {output_root.as_posix()}", quiet=quiet)
    _log(f"Style:            {resolved_style}\n", quiet=quiet)

    for source in files:
        relative = source.relative_to(input_root)
        destination = output_root / relative.with_suffix(".docx")
        if destination.exists() and not force and not dry_run:
            _error(
                "output file already exists.",
                code=EXIT_OUTPUT_PATH,
                path=destination,
                hint="Use --force to overwrite existing files.",
            )

        report = validate_markdown_file(
            source,
            strict=False,
            check_pii=bool(validation_cfg.get("check_pii", False)),
            check_placeholders=bool(validation_cfg.get("check_placeholders", True)),
            check_tables=bool(validation_cfg.get("check_empty_tables", True)),
            check_headings=bool(validation_cfg.get("check_heading_order", True)),
            check_links=bool(validation_cfg.get("check_links", False)),
        )
        reports.append(report)
        if report.errors or (strict_mode and report.warnings):
            _print_report(report, quiet=quiet, verbose=True)
            _error("validation failed during batch conversion.", code=EXIT_VALIDATION, path=source)

        if dry_run:
            _log(f"[dry-run] {source.as_posix()} -> {destination.as_posix()}", quiet=quiet)
            continue

        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            convert_markdown_to_docx(
                source,
                destination,
                style=resolved_style,
                template_path=resolved_template,
            )
        except Exception as exc:  # pragma: no cover - backend/runtime failure
            _error(f"conversion backend error: {exc}", code=EXIT_BACKEND, path=source)
        converted += 1
        if verbose and not quiet:
            print(f"[mdx] converted {source.as_posix()} -> {destination.as_posix()}")

    warnings, errors = summarize_reports(reports)
    _log(f"\nFiles discovered: {len(files)}", quiet=quiet)
    _log(f"Files converted:  {converted if not dry_run else 0}", quiet=quiet)
    _log(f"Warnings:         {warnings}", quiet=quiet)
    _log(f"Errors:           {errors}", quiet=quiet)
    _log(f"\nStatus: {'dry-run complete' if dry_run else 'complete'}", quiet=quiet)
    return None


@registry.register(
    "validate",
    description="Validate Markdown files before conversion.",
    examples=["mdx validate docs --strict --check_placeholders --check_pii"],
    render=False,
)
@registry.argument("input_path", type=t.Path(exists=True), help="Markdown file or directory.")
@registry.argument("strict", type=bool, default=False, help="Treat warnings as errors.")
@registry.argument("check_pii", type=bool, default=False, help="Detect likely PII.")
@registry.argument("check_placeholders", type=bool, default=False, help="Detect placeholder text.")
@registry.argument("check_tables", type=bool, default=False, help="Validate table rows/cells.")
@registry.argument("check_headings", type=bool, default=False, help="Validate heading hierarchy.")
@registry.argument("check_links", type=bool, default=False, help="Validate local links.")
@registry.argument("verbose", type=bool, default=False, help="Detailed validation output.")
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def validate(
    input_path: Path,
    strict: bool = False,
    check_pii: bool = False,
    check_placeholders: bool = False,
    check_tables: bool = False,
    check_headings: bool = False,
    check_links: bool = False,
    verbose: bool = False,
    quiet: bool = False,
) -> None:
    target = Path(input_path)
    reports: list[ValidationReport] = []

    checks = {
        "check_pii": check_pii,
        "check_placeholders": check_placeholders,
        "check_tables": check_tables,
        "check_headings": check_headings,
        "check_links": check_links,
    }
    if not any(checks.values()):
        checks = {
            "check_pii": False,
            "check_placeholders": True,
            "check_tables": True,
            "check_headings": True,
            "check_links": False,
        }

    if target.is_file():
        reports.append(
            validate_markdown_file(
                target,
                strict=strict,
                check_pii=checks["check_pii"],
                check_placeholders=checks["check_placeholders"],
                check_tables=checks["check_tables"],
                check_headings=checks["check_headings"],
                check_links=checks["check_links"],
            )
        )
    else:
        files = _discover_markdown_files(target, include="*.md", exclude=None)
        for file_path in files:
            reports.append(
                validate_markdown_file(
                    file_path,
                    strict=strict,
                    check_pii=checks["check_pii"],
                    check_placeholders=checks["check_placeholders"],
                    check_tables=checks["check_tables"],
                    check_headings=checks["check_headings"],
                    check_links=checks["check_links"],
                )
            )

    for report in reports:
        _print_report(report, quiet=quiet, verbose=verbose)
        if not quiet and len(reports) > 1:
            print()

    warnings, errors = summarize_reports(reports)
    if errors:
        _error("validation failed.", code=EXIT_VALIDATION)
    if strict and warnings:
        _error("validation warnings found in strict mode.", code=EXIT_VALIDATION)

    if not quiet:
        print(f"Summary: warnings={warnings}, errors={errors}")
    return None


@registry.register(
    "preview",
    description="Generate an HTML preview from Markdown.",
    examples=["mdx preview docs/example.md --output dist/preview/example.html"],
    render=False,
)
@registry.argument("input_md", type=t.Path(exists=True, readable=True), help="Input Markdown file.")
@registry.argument("output", type=str, help="Output HTML path.")
@registry.argument("style", type=str, default="", help="Style preset.")
@registry.argument("config", type=str, default="mdx.toml", help="Config file path.")
@registry.argument("template", type=str, default="", help="Reserved for compatibility.")
@registry.argument("force", type=bool, default=False, help="Overwrite output.")
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def preview(
    input_md: Path,
    output: str,
    style: str = "",
    config: str = "mdx.toml",
    template: str = "",
    force: bool = False,
    quiet: bool = False,
) -> None:
    source = Path(input_md)
    destination = Path(output)
    _ensure_input_file(source)
    _ensure_write_target(destination, force=force)

    config_dict = _resolve_config(config, explicit=(config != "mdx.toml"))
    resolved_style = _resolve_style(style, config_dict)
    _ = template  # reserved for future template-aware preview renderers
    css = _html_css_for_style(resolved_style)
    html_body = render_markdown_to_html(source.read_text(encoding="utf-8"))
    html = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>MDX Preview</title>
  <style>{css}</style>
</head>
<body>
  <main class="content">
{html_body}
  </main>
</body>
</html>
"""
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(html, encoding="utf-8")
    _log("Preview generated:", quiet=quiet)
    _log(destination.as_posix(), quiet=quiet)
    return None


def _html_css_for_style(style: str) -> str:
    if style == "academic":
        return "body{font-family:'Times New Roman',serif;background:#fff;color:#111;line-height:1.6;} .content{max-width:900px;margin:2rem auto;padding:0 1rem;} pre{background:#f8f8f8;padding:.8rem;overflow:auto;}"
    if style == "technical":
        return "body{font-family:'Segoe UI',sans-serif;background:#fff;color:#111;line-height:1.55;} .content{max-width:960px;margin:2rem auto;padding:0 1rem;} code,pre{font-family:Consolas,monospace;background:#f6f8fa;} pre{padding:.8rem;overflow:auto;}"
    if style == "executive":
        return "body{font-family:Cambria,serif;background:#fff;color:#1a1a1a;line-height:1.65;} .content{max-width:920px;margin:2rem auto;padding:0 1rem;} pre{background:#f8f8f8;padding:.8rem;overflow:auto;}"
    if style == "minimal":
        return "body{font-family:Arial,sans-serif;background:#fff;color:#111;line-height:1.6;} .content{max-width:900px;margin:2rem auto;padding:0 1rem;} pre{background:#f4f4f4;padding:.8rem;overflow:auto;}"
    return "body{font-family:Calibri,sans-serif;background:#fff;color:#111;line-height:1.6;} .content{max-width:920px;margin:2rem auto;padding:0 1rem;} pre{background:#f6f8fa;padding:.8rem;overflow:auto;} table{border-collapse:collapse;} td,th{border:1px solid #d0d7de;padding:6px 8px;}"


@registry.register("doctor", description="Check local environment readiness for MDX.", render=False)
@registry.argument("quiet", type=bool, default=False, help="Suppress non-error output.")
def doctor(quiet: bool = False) -> None:
    checks: list[tuple[str, bool, str]] = []

    runtime_ok = sys.version_info >= (3, 11)
    checks.append(("Runtime supported", runtime_ok, f"Python {sys.version.split()[0]}"))

    backend_ok = True
    backend_msg = "python-docx and markdown-it-py available"
    try:
        import docx  # noqa: F401
        import markdown_it  # noqa: F401
    except Exception as exc:  # pragma: no cover - import/runtime failure
        backend_ok = False
        backend_msg = str(exc)
    checks.append(("DOCX backend available", backend_ok, backend_msg))

    config_path = Path("mdx.toml")
    checks.append(("Config file readable", config_path.exists(), config_path.as_posix()))

    template_path = Path("templates/professional.docx")
    checks.append(("Template path exists", template_path.exists(), template_path.as_posix()))

    out_dir = Path("docs")
    writable = True
    detail = out_dir.as_posix()
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
        probe = out_dir / ".mdx-write-test"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
    except Exception as exc:  # pragma: no cover - fs/runtime failure
        writable = False
        detail = str(exc)
    checks.append(("Output directory writable", writable, detail))

    mermaid_available = shutil.which("mmdc") is not None
    checks.append(("Mermaid renderer available", mermaid_available, "mmdc on PATH"))

    pandoc_available = shutil.which("pandoc") is not None
    checks.append(("Pandoc availability", pandoc_available, "pandoc on PATH"))

    if not quiet:
        print("MDX Doctor\n")
        for label, ok, detail in checks:
            mark = "[OK]" if ok else "[FAIL]"
            print(f"{mark} {label}")
            if detail:
                print(f"  {detail}")

    if all(ok for _, ok, _ in checks[:2]):
        _log("\nStatus: ready", quiet=quiet)
        return None

    _error("environment is not ready.", code=EXIT_BACKEND)
    return None


@registry.register("version", description="Show installed version.")
def version() -> str:
    return f"mdx {__version__}"


def main() -> None:
    try:
        registry.run(shell_title="MDX CLI", shell_usage=True)
    except SystemExit:
        raise
    except Exception as exc:  # pragma: no cover - catch-all guard
        _error(str(exc), code=EXIT_GENERAL)


if __name__ == "__main__":
    main()
